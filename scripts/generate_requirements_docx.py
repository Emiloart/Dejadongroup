from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "De-Jadon-Group-Full-Website-Requirements.docx"

ORANGE = "E2691B"
GREEN = "4F7A3D"
INK = "241F19"
WARM = "FAF5EC"
LIGHT_ORANGE = "FCE8D6"
LIGHT_GREEN = "E7F0DF"
LIGHT_STONE = "F6F2EA"
WHITE = "FFFFFF"
MUTED = "625A50"

REQUIRED = "Required from De Jadon Group. Do not publish until supplied and approved."
APPROVAL = "Client approval required before final implementation."
LEGAL = "Legal/commercial approval required before publication."
MEDIA = "Client to supply final asset, caption, ownership confirmation, and alt text."


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "E9DED0") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_begin)
    run._r.append(instr_text)
    run._r.append(fld_end)


def style_document(doc: Document) -> None:
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in [
        ("Title", 26, INK),
        ("Heading 1", 18, INK),
        ("Heading 2", 14, INK),
        ("Heading 3", 11, ORANGE),
    ]:
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)

    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)
        section.header_distance = Inches(0.25)
        section.footer_distance = Inches(0.25)

        header = section.header.paragraphs[0]
        header.text = "De Jadon Group Website Requirements"
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

        footer = section.footer.paragraphs[0]
        add_page_number(footer)


def add_rule(doc: Document, color: str = ORANGE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    set_cell_shading(cell, color)
    cell.height = Inches(0.03)
    cell.text = ""


def add_note(doc: Document, title: str, body: str, fill: str = LIGHT_ORANGE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    title_run = paragraph.add_run(f"{title}: ")
    title_run.bold = True
    title_run.font.color.rgb = RGBColor.from_string(INK)
    body_run = paragraph.add_run(body)
    body_run.font.color.rgb = RGBColor.from_string(INK)


def add_heading(doc: Document, text: str, level: int = 1, accent: str | None = None) -> None:
    heading = doc.add_heading(text, level=level)
    if accent:
        heading.runs[0].font.color.rgb = RGBColor.from_string(accent)


def add_para(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        run.bold = True
        paragraph.add_run(text[len(bold_prefix) :])
    else:
        paragraph.add_run(text)


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.add_run(item)


def add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        cell = header_cells[index]
        set_cell_shading(cell, ORANGE)
        set_cell_border(cell, "D85F16")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = Inches(widths[index])
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
        run.font.size = Pt(8.5)

    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        fill = WHITE if row_index % 2 == 0 else LIGHT_STONE
        for cell_index, value in enumerate(row):
            cell = cells[cell_index]
            set_cell_shading(cell, fill)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if widths:
                cell.width = Inches(widths[cell_index])
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.add_run(str(value))

    doc.add_paragraph()


def add_requirement_table(doc: Document, rows: Sequence[Sequence[str]]) -> None:
    add_table(
        doc,
        ["Ref", "Required information or content", "Format / validation", "Owner / status"],
        rows,
        [0.75, 3.0, 2.25, 1.85],
    )


def add_asset_table(doc: Document, rows: Sequence[Sequence[str]]) -> None:
    add_table(
        doc,
        ["Asset area", "Required asset", "Specification", "Approval note"],
        rows,
        [1.8, 2.2, 2.5, 1.4],
    )


def add_section_break(doc: Document, title: str) -> None:
    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, title, 1)
    add_rule(doc)


def content_slot_rows(prefix: str, slots: Sequence[str]) -> list[list[str]]:
    return [
        [
            f"{prefix}-{index:02d}",
            slot,
            "Final approved heading, body copy, CTA label, CTA target, image requirement if applicable, SEO note, and mobile priority.",
            REQUIRED,
        ]
        for index, slot in enumerate(slots, start=1)
    ]


def build_document() -> Document:
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("De Jadon Group")
    run.bold = True
    run.font.size = Pt(30)
    run.font.color.rgb = RGBColor.from_string(INK)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Full Website Requirements and Content Collection Document")
    subtitle_run.bold = True
    subtitle_run.font.size = Pt(16)
    subtitle_run.font.color.rgb = RGBColor.from_string(ORANGE)

    add_rule(doc)
    add_note(
        doc,
        "Content rule",
        "No generic wording, placeholder business claims, estimated pricing, fake descriptions, fake statistics, or assumed company information should go live. Each live detail must be supplied or approved by De Jadon Group.",
    )

    add_para(
        doc,
        "This document converts the current frontend skeleton into a complete content and requirements checklist for the final De Jadon Group website. It covers every public route, visible section, dashboard shell, image canvas, interaction, form, and approval point currently represented in the site.",
    )
    add_para(
        doc,
        "The current site is a frontend skeleton only. It intentionally contains named slots and image canvases instead of final copy. The purpose of this file is to collect the exact information needed before those slots are replaced with client-approved content.",
    )

    add_heading(doc, "Document Control", 2)
    add_table(
        doc,
        ["Item", "Detail"],
        [
            ["Project", "De Jadon Group website and portal frontend"],
            ["Current active businesses", "Agriculture and Agro-Real Estate"],
            ["Excluded from public navigation for now", "Real Estate, Construction, African Tales"],
            ["Dashboard exposure", "Partner Dashboard and Client Dashboard only. Admin dashboard must not appear in user-facing navigation or demo buttons."],
            ["Frontend stack", "Next.js App Router, TypeScript, Tailwind CSS, lucide-react"],
            ["Backend status", "Not implemented in this phase. Shared identity, investments, referrals, wallet, commissions, and admin management are future integrations."],
        ],
        [2.2, 5.45],
    )

    add_heading(doc, "How To Use This Document", 2)
    add_numbered(
        doc,
        [
            "De Jadon Group supplies exact approved content for every required item marked as required.",
            "Designer or developer replaces matching frontend content slots and image canvases with approved content only.",
            "Commercial, investment, partner compensation, legal, and payment details must be approved by the correct business owner before publication.",
            "After content insertion, every route must be reviewed on desktop and mobile before launch.",
        ],
    )

    add_section_break(doc, "1. Scope and Exclusions")
    add_heading(doc, "In Scope", 2)
    add_bullets(
        doc,
        [
            "Public marketing website: Home, About, Businesses, Agriculture, Agro-Real Estate, Partner Program, Contact, and Login.",
            "Business pages for Agriculture and Agro-Real Estate only.",
            "Agriculture product catalog structure that can later grow into marketplace functionality.",
            "Demo Partner Dashboard and Client Dashboard shells with named areas only.",
            "Responsive layout requirements for desktop, tablet, and mobile.",
            "Content collection requirements for every named slot, image canvas, form, and dashboard block.",
        ],
    )
    add_heading(doc, "Out Of Scope For This Phase", 2)
    add_bullets(
        doc,
        [
            "No backend, database, NestJS API, Prisma schema, PostgreSQL setup, or authentication logic.",
            "No live payment, wallet, commission, withdrawal, KYC, document upload, referral tree calculation, order checkout, or marketplace transaction logic.",
            "No public Real Estate, Construction, or African Tales pages.",
            "No user-facing Admin Dashboard, admin demo button, admin route link, or admin navigation entry.",
            "No invented prices, claims, ROI, delivery terms, farm sizes, locations, CEO biography, testimonials, product descriptions, or statistics.",
        ],
    )

    add_heading(doc, "Core Product Positioning", 2)
    add_para(
        doc,
        "De Jadon Group should be presented as the parent company. Agriculture and Agro-Real Estate are the active businesses currently shown on the website. The platform should be flexible enough to add future businesses later without redesigning the entire site.",
    )

    add_section_break(doc, "2. Site Architecture")
    add_heading(doc, "Primary Navigation", 2)
    add_table(
        doc,
        ["Navigation item", "Route", "Purpose", "Content required before launch"],
        [
            ["Home", "/", "Main entry point and overview of active businesses.", "Final hero copy, about preview, CEO details, CTA, updates, FAQs, and images."],
            ["About", "/about", "Company identity, vision, mission, why choose us, and CEO profile.", "Official company story, vision, mission, positioning, CEO profile, and approved images."],
            ["Businesses", "/businesses", "Index of active businesses only.", "Overview copy and approved cards for Agriculture and Agro-Real Estate."],
            ["Partner Program", "/partner-program", "Partner registration and partner growth program overview.", "Program rules, tiers, benefits, compensation, FAQ, terms, and register/login messaging."],
            ["Contact", "/contact", "Inquiry routing and contact form.", "Phone, email, office address, map/location, working hours, and inquiry handling rules."],
            ["Login", "/login", "Frontend access point for partner and client portal demos.", "Final portal copy, shared identity rules, login/register/forgot-password wording, and access states."],
        ],
        [1.25, 1.45, 2.3, 2.65],
    )

    add_heading(doc, "Route Matrix", 2)
    add_table(
        doc,
        ["Area", "Route", "Current frontend status", "Required next decision"],
        [
            ["Home", "/", "Implemented as skeleton with hero carousel and named content slots.", "Supply approved content and images for every slot."],
            ["About", "/about", "Implemented as skeleton.", "Supply official company identity and CEO details."],
            ["Businesses", "/businesses", "Shows only Agriculture and Agro-Real Estate.", "Confirm no additional business appears before launch."],
            ["Agriculture", "/businesses/agriculture", "Implemented as active business with green accents.", "Supply Agriculture overview, sections, and product details."],
            ["Crop Farming", "/businesses/agriculture/crop-farming", "Shows Palm and Pepper cards.", "Supply crop-specific content and imagery."],
            ["Livestock", "/businesses/agriculture/livestock", "Shows Poultry, Fishery, and BSF.", "Supply livestock category details and imagery."],
            ["Poultry", "/businesses/agriculture/livestock/poultry", "Skeleton detail page.", "Supply poultry details and gallery assets."],
            ["Fishery", "/businesses/agriculture/livestock/fishery", "Skeleton detail page.", "Supply fishery details and gallery assets."],
            ["BSF", "/businesses/agriculture/livestock/bsf", "Skeleton detail page.", "Supply BSF details and gallery assets."],
            ["Agriculture Products", "/businesses/agriculture/products", "Shows named product categories only.", "Supply final product catalog fields and images."],
            ["Agriculture Gallery", "/businesses/agriculture/gallery", "Three gallery canvases.", "Supply final gallery assets."],
            ["Agro-Real Estate", "/businesses/agro-real-estate", "Implemented as separate business.", "Supply real estate investment content and legal details."],
            ["Buy Cultivated Farmland", "/businesses/agro-real-estate/buy-cultivated-farmland", "Skeleton detail and inquiry blocks.", "Supply offer details, land documents, process, and terms."],
            ["Buy Land & Managed Cultivation", "/businesses/agro-real-estate/managed-cultivation", "Skeleton detail and inquiry blocks.", "Supply managed cultivation process and obligations."],
            ["Investment Plans", "/businesses/agro-real-estate/investment-plans", "Skeleton detail and inquiry blocks.", "Supply approved plan names, inclusions, terms, risks, and pricing if applicable."],
            ["Agro-Real Estate FAQs", "/businesses/agro-real-estate/faqs", "Skeleton detail and inquiry blocks.", "Supply approved question and answer set."],
            ["Agro-Real Estate Gallery", "/businesses/agro-real-estate/gallery", "Three gallery canvases.", "Supply final gallery assets."],
            ["Contact", "/contact", "Inquiry cards plus form.", "Supply routing and contact handling rules."],
            ["Login", "/login", "Partner and Client demo buttons only.", "Define final auth wording and account states."],
            ["Partner Dashboard", "/dashboard/partner", "User-facing demo shell.", "Supply all required field labels, empty states, and future data rules."],
            ["Client Dashboard", "/dashboard/client", "User-facing demo shell.", "Supply all required field labels, empty states, and future data rules."],
        ],
        [1.55, 2.15, 2.1, 2.05],
    )

    add_section_break(doc, "3. Brand, Content, and Asset Standards")
    add_heading(doc, "Brand Requirements", 2)
    add_requirement_table(
        doc,
        [
            ["BR-01", "Official legal name and preferred public display name.", "Text exactly as De Jadon Group wants it displayed.", REQUIRED],
            ["BR-02", "Logo files and usage rules.", "SVG, PNG, transparent PNG, light version, dark version, favicon, and minimum clear-space rules.", MEDIA],
            ["BR-03", "Approved brand colors.", "Confirm current white, warm white, orange, and agriculture green values or provide updated color codes.", APPROVAL],
            ["BR-04", "Typography preference.", "Confirm current frontend fonts or provide brand font files and licenses.", APPROVAL],
            ["BR-05", "Tone of voice.", "Formal, professional, investor-facing, partner-facing, or sales-focused tone must be defined by the client.", REQUIRED],
            ["BR-06", "Legal footer details.", "Registered company details, copyright wording, privacy policy link, terms link, and any required disclaimers.", LEGAL],
            ["BR-07", "Social links.", "Official handles for Facebook, Instagram, YouTube, TikTok, LinkedIn, X, WhatsApp, and any other approved platform.", REQUIRED],
        ]
    )

    add_heading(doc, "Content Rules", 2)
    add_bullets(
        doc,
        [
            "Every public claim must have a source supplied by De Jadon Group.",
            "No yield, ROI, harvest, commission, land-size, income, or timeline claim should be written without approval.",
            "Every image must be owned by De Jadon Group or licensed for commercial website use.",
            "Every person shown in photos must have permission for website use where required.",
            "Every CTA must have a defined target route, form, phone number, WhatsApp link, or portal action.",
        ],
    )

    add_heading(doc, "Global Image Canvas Requirements", 2)
    add_asset_table(
        doc,
        [
            ["Hero canvases", "Main visual for each hero slide and page hero.", "Landscape image. Subject must be clear and relevant. Avoid dark, cropped, or generic stock visuals.", MEDIA],
            ["Business cards", "Agriculture and Agro-Real Estate image canvases.", "One strong image per business. Should visually communicate the exact business area.", MEDIA],
            ["Gallery canvases", "Three or more images per gallery page.", "Each image requires caption, location if displayed, date if relevant, and alt text.", MEDIA],
            ["CEO image", "CEO portrait.", "Professional portrait with approved name, title, and biography.", MEDIA],
            ["Product images", "One image per product category and later one image per product item.", "Clean product photos with consistent crop and approved captions.", MEDIA],
            ["Dashboard canvases", "QR code, referral tree, photos, reports, and marketing-material previews.", "Use live-generated data later. Current frontend should show only structural placeholders.", APPROVAL],
        ],
    )

    add_section_break(doc, "4. Global Header and Footer")
    add_heading(doc, "Header Requirements", 2)
    add_requirement_table(
        doc,
        [
            ["HD-01", "Logo lockup content.", "Confirm whether the `DJ` mark remains, whether a full logo replaces it, and what exact brand text appears beside it.", REQUIRED],
            ["HD-02", "Primary navigation.", "Home, About, Businesses, Partner Program, Contact, Login. No Coming Soon businesses and no Admin Dashboard.", APPROVAL],
            ["HD-03", "Header CTA buttons.", "Confirm whether Partner Program and Contact remain as desktop CTAs and define final labels.", APPROVAL],
            ["HD-04", "Mobile navigation.", "All primary nav items must be reachable behind the mobile menu and close after navigation.", APPROVAL],
            ["HD-05", "Active states.", "Define whether current page indicator is required for launch.", APPROVAL],
        ]
    )

    add_heading(doc, "Footer Requirements", 2)
    add_requirement_table(
        doc,
        [
            ["FT-01", "Footer company block.", "Final footer description or leave minimal with brand only.", REQUIRED],
            ["FT-02", "Footer navigation.", "Same public routes as header. Admin and excluded businesses must not be present.", APPROVAL],
            ["FT-03", "Footer businesses.", "Agriculture and Agro-Real Estate only.", APPROVAL],
            ["FT-04", "Contact footer area.", "Final contact label, phone/email/WhatsApp if displayed, and route target.", REQUIRED],
            ["FT-05", "Legal links.", "Privacy Policy, Terms, Partner Terms, Investment Terms, and Refund/Cancellation Policy if applicable.", LEGAL],
        ]
    )

    add_section_break(doc, "5. Home Page Requirements")
    add_heading(doc, "Home Page Structure", 2)
    add_bullets(
        doc,
        [
            "Hero carousel with four slots: De Jadon Group, Agriculture, Agro-Real Estate, Partner Program.",
            "About De Jadon Group preview section.",
            "Businesses carousel showing Agriculture and Agro-Real Estate.",
            "Partner Program section below the business carousel.",
            "CEO section.",
            "Call To Action section.",
            "News & Updates section.",
            "FAQ section.",
        ],
    )

    add_heading(doc, "Hero Carousel Content", 2)
    add_requirement_table(
        doc,
        [
            ["HM-HR-01", "De Jadon Group Intro Slot.", "Final headline support copy, image, primary CTA target, secondary CTA target, and mobile crop instructions.", REQUIRED],
            ["HM-HR-02", "Agriculture Hero Slot.", "Approved agriculture copy, Agriculture Hero Canvas image, View Agriculture CTA, View Products CTA.", REQUIRED],
            ["HM-HR-03", "Agro-Real Estate Hero Slot.", "Approved Agro-Real Estate copy, image, Investment Plans CTA, Contact CTA.", REQUIRED],
            ["HM-HR-04", "Partner Program Hero Slot.", "Approved partner-program copy, image, Open Partner Program CTA, Login CTA.", REQUIRED],
            ["HM-HR-05", "Carousel behavior.", "Autoplay interval, manual arrows, dots, mobile swipe if added later, pause behavior, and reduced-motion behavior.", APPROVAL],
        ]
    )

    add_heading(doc, "Home Content Slots", 2)
    add_requirement_table(
        doc,
        content_slot_rows(
            "HM",
            [
                "About De Jadon Group Content Slot",
                "About Image Canvas",
                "Businesses carousel content for Agriculture",
                "Businesses carousel content for Agro-Real Estate",
                "Partner Program Homepage Slot",
                "CEO Profile Slot",
                "CEO Image Canvas",
                "De Jadon Group CTA Slot",
                "Featured Update Slot",
                "Featured Update Image Canvas",
                "Agriculture Update Slot",
                "Agro-Real Estate Update Slot",
                "Partner Program Update Slot",
                "Agriculture FAQ",
                "Agro-Real Estate FAQ",
                "Partner Program FAQ",
                "Contact FAQ",
            ],
        )
    )

    add_section_break(doc, "6. About Page Requirements")
    add_heading(doc, "About Page Content Slots", 2)
    add_requirement_table(
        doc,
        content_slot_rows(
            "AB",
            [
                "About De Jadon Group Intro Slot",
                "About Image Canvas",
                "About Content Slot",
                "Vision Content Slot",
                "Mission Content Slot",
                "Why Choose Us Content Slot",
                "CEO Profile Slot",
                "CEO Image Canvas",
            ],
        )
    )
    add_heading(doc, "CEO Requirements", 2)
    add_requirement_table(
        doc,
        [
            ["CEO-01", "CEO full name.", "Exact spelling and title.", REQUIRED],
            ["CEO-02", "CEO role description.", "Approved short and long version.", REQUIRED],
            ["CEO-03", "CEO biography.", "Approved biography only. No invented background, awards, or story.", REQUIRED],
            ["CEO-04", "CEO image.", "Approved portrait, caption, image rights, and alt text.", MEDIA],
            ["CEO-05", "CEO quote if used.", "Exact quote approved by CEO or authorized representative.", APPROVAL],
        ]
    )

    add_section_break(doc, "7. Businesses Index Requirements")
    add_heading(doc, "Businesses Page", 2)
    add_requirement_table(
        doc,
        [
            ["BS-01", "Businesses Overview Slot.", "Short parent-company explanation that introduces active businesses only.", REQUIRED],
            ["BS-02", "Businesses Image Canvas.", "Approved image representing De Jadon Group or the active business ecosystem.", MEDIA],
            ["BS-03", "Agriculture business card.", "Approved card copy, image, CTA label, and route.", REQUIRED],
            ["BS-04", "Agro-Real Estate business card.", "Approved card copy, image, CTA label, and route.", REQUIRED],
            ["BS-05", "Excluded business rule.", "Real Estate, Construction, and African Tales must not appear in this frontend until client approves their launch.", APPROVAL],
        ]
    )

    add_section_break(doc, "8. Agriculture Requirements")
    add_heading(doc, "Agriculture Overview", 2)
    add_requirement_table(
        doc,
        content_slot_rows(
            "AG",
            [
                "Agriculture Overview Slot",
                "Agriculture Image Canvas",
                "Agriculture Sections intro if needed",
                "Products summary section",
            ],
        )
    )

    add_heading(doc, "Crop Farming", 2)
    add_requirement_table(
        doc,
        [
            ["AG-CF-01", "Crop Farming Overview Slot.", "Exact wording for crop farming activities and scope.", REQUIRED],
            ["AG-CF-02", "Palm.", "Approved palm farming details, image, availability, location if displayed, and CTA.", REQUIRED],
            ["AG-CF-03", "Pepper.", "Approved pepper farming details, image, availability, location if displayed, and CTA.", REQUIRED],
            ["AG-CF-04", "Farming method claims.", "Any claims about practices, quality, scale, harvest period, or output must be supplied and approved.", APPROVAL],
        ]
    )

    add_heading(doc, "Livestock", 2, GREEN)
    add_requirement_table(
        doc,
        [
            ["AG-LS-01", "Livestock Overview Slot.", "Exact wording for livestock operations under Agriculture.", REQUIRED],
            ["AG-LS-02", "Poultry Detail Slot.", "Approved poultry details, operating scope, images, product relation if any, and CTA.", REQUIRED],
            ["AG-LS-03", "Poultry Gallery Canvas.", "Approved poultry gallery images and captions.", MEDIA],
            ["AG-LS-04", "Fishery Detail Slot.", "Approved fishery details, operating scope, images, product relation if any, and CTA.", REQUIRED],
            ["AG-LS-05", "Fishery Gallery Canvas.", "Approved fishery gallery images and captions.", MEDIA],
            ["AG-LS-06", "BSF (Black Soldier Fly) Detail Slot.", "Approved BSF details, use case, operating scope, images, and CTA.", REQUIRED],
            ["AG-LS-07", "BSF (Black Soldier Fly) Gallery Canvas.", "Approved BSF gallery images and captions.", MEDIA],
        ]
    )

    add_heading(doc, "Agriculture Products", 2, GREEN)
    add_para(
        doc,
        "Agriculture Products remain under Agriculture. The future marketplace should grow from this section, but the current phase only requires a product catalog structure.",
    )
    add_requirement_table(
        doc,
        [
            ["AG-PR-01", "Live Fish.", "Approved product category copy, product image, ordering instruction, availability rule, and CTA.", REQUIRED],
            ["AG-PR-02", "Dried/Packaged Fish.", "Approved product category copy, packaging details, image, ordering instruction, availability rule, and CTA.", REQUIRED],
            ["AG-PR-03", "Palm.", "Approved product category copy, image, ordering instruction, availability rule, and CTA.", REQUIRED],
            ["AG-PR-04", "Pepper.", "Approved product category copy, image, ordering instruction, availability rule, and CTA.", REQUIRED],
            ["AG-PR-05", "Other Farm Harvests.", "Approved category scope, image, ordering instruction, availability rule, and CTA.", REQUIRED],
            ["AG-PR-06", "Future marketplace expansion.", "Define future cart, checkout, delivery, pickup, payment, inventory, wholesale, and order-tracking requirements before build.", APPROVAL],
        ]
    )

    add_heading(doc, "Agriculture Gallery", 2, GREEN)
    add_asset_table(
        doc,
        [
            ["Agriculture Gallery Canvas 1", "Approved Agriculture gallery image.", "Image, caption, alt text, ownership confirmation.", MEDIA],
            ["Agriculture Gallery Canvas 2", "Approved Agriculture gallery image.", "Image, caption, alt text, ownership confirmation.", MEDIA],
            ["Agriculture Gallery Canvas 3", "Approved Agriculture gallery image.", "Image, caption, alt text, ownership confirmation.", MEDIA],
        ]
    )

    add_section_break(doc, "9. Agro-Real Estate Requirements")
    add_heading(doc, "Positioning", 2)
    add_para(
        doc,
        "Agro-Real Estate is treated as its own business, separate from Agriculture. It covers farmland purchase, cultivated farmland, managed cultivation, investment plans, FAQs, and gallery content.",
    )

    add_heading(doc, "Agro-Real Estate Content Slots", 2)
    add_requirement_table(
        doc,
        content_slot_rows(
            "ARE",
            [
                "Agro-Real Estate Overview Slot",
                "Agro-Real Estate Image Canvas",
                "Buy Cultivated Farmland Overview Slot",
                "Buy Cultivated Farmland Details Slot",
                "Buy Cultivated Farmland Inquiry Slot",
                "Buy Land & Managed Cultivation Overview Slot",
                "Buy Land & Managed Cultivation Details Slot",
                "Buy Land & Managed Cultivation Inquiry Slot",
                "Investment Plans Overview Slot",
                "Investment Plans Details Slot",
                "Investment Plans Inquiry Slot",
                "FAQs Overview Slot",
                "FAQs Details Slot",
                "FAQs Inquiry Slot",
            ],
        )
    )

    add_heading(doc, "Commercial and Legal Information", 2)
    add_requirement_table(
        doc,
        [
            ["ARE-LG-01", "Land location information.", "Exact state, LGA, community, estate/farm name, coordinates if displayed, and map rules.", LEGAL],
            ["ARE-LG-02", "Land title and documentation.", "Approved description of title documents, verification process, deed/C of O/survey status if applicable.", LEGAL],
            ["ARE-LG-03", "Land size units.", "Define plot, acre, hectare, or other units and conversion rules.", APPROVAL],
            ["ARE-LG-04", "Payment terms.", "Approved payment schedule, deposit rules, installment rules, receipts, refund/cancellation terms.", LEGAL],
            ["ARE-LG-05", "Managed cultivation obligations.", "Who cultivates, what is cultivated, management duration, maintenance scope, reporting obligations, and client responsibilities.", LEGAL],
            ["ARE-LG-06", "Investment plan terms.", "Approved plan names, inclusions, exclusions, eligibility, duration, and payment requirements.", LEGAL],
            ["ARE-LG-07", "Returns, yield, or ROI language.", "Do not publish any return/yield/ROI claim unless exact wording and basis are supplied and legally approved.", LEGAL],
            ["ARE-LG-08", "Risk disclosures.", "Agricultural, weather, market, operational, title, and payment risks if required by legal review.", LEGAL],
            ["ARE-LG-09", "Client dashboard data.", "Define land records, documents, farm stage, photos, reports, and yield history data sources.", APPROVAL],
        ]
    )

    add_heading(doc, "Agro-Real Estate Gallery", 2)
    add_asset_table(
        doc,
        [
            ["Agro-Real Estate Gallery Canvas 1", "Approved gallery image.", "Image, caption, alt text, ownership confirmation.", MEDIA],
            ["Agro-Real Estate Gallery Canvas 2", "Approved gallery image.", "Image, caption, alt text, ownership confirmation.", MEDIA],
            ["Agro-Real Estate Gallery Canvas 3", "Approved gallery image.", "Image, caption, alt text, ownership confirmation.", MEDIA],
        ]
    )

    add_section_break(doc, "10. Partner Program Requirements")
    add_heading(doc, "Partner Program Structure", 2)
    add_requirement_table(
        doc,
        content_slot_rows(
            "PP",
            [
                "Partner Program Overview Slot",
                "Partner Program Image Canvas",
                "Overview Section Slot",
                "Benefits Section Slot",
                "Partner Tiers Section Slot",
                "Compensation Section Slot",
                "FAQs Section Slot",
                "Register Section Slot",
                "Login Section Slot",
                "Partner Program Homepage Slot",
            ],
        )
    )

    add_heading(doc, "Partner Tiers", 2)
    add_requirement_table(
        doc,
        [
            ["PP-TI-01", "Tier names.", "Exact approved tier names. Use `Partner Tier` terminology consistently, not `Package`.", REQUIRED],
            ["PP-TI-02", "Tier fees if any.", "Exact approved amount, currency, payment rules, and receipt process. Leave blank if no fee.", LEGAL],
            ["PP-TI-03", "Tier benefits.", "Approved benefits per tier. No assumed benefits or fake privileges.", APPROVAL],
            ["PP-TI-04", "Tier qualification.", "Approval process, KYC requirements, renewal, upgrade, downgrade, and suspension rules.", LEGAL],
            ["PP-TI-05", "Product access rules.", "Define which products partners can sell, if tier-based product access exists.", APPROVAL],
        ]
    )

    add_heading(doc, "Compensation and Referral Rules", 2)
    add_requirement_table(
        doc,
        [
            ["PP-CP-01", "Direct referral rule.", "Approved commission amount or percentage, trigger event, and eligibility.", LEGAL],
            ["PP-CP-02", "Downline rule.", "Approved depth, genealogy logic, team volume logic, and cap if any.", LEGAL],
            ["PP-CP-03", "Sales commission rule.", "Approved commission by product or sale type, payment timing, reversal rules, and refund handling.", LEGAL],
            ["PP-CP-04", "Leadership or bonus rule.", "Approved qualification criteria, payout basis, and frequency.", LEGAL],
            ["PP-CP-05", "Wallet rule.", "Balance types, pending/available state, withdrawal threshold, bank verification, and payout approval workflow.", LEGAL],
            ["PP-CP-06", "Partner conduct.", "Approved terms preventing false claims, unauthorized pricing, misleading investment promises, or misuse of De Jadon Group branding.", LEGAL],
        ]
    )

    add_heading(doc, "Partner Registration", 2)
    add_requirement_table(
        doc,
        [
            ["PP-RG-01", "Personal information.", "Full name, email, phone, address, passport photo, ID type, ID number, and other KYC fields if required.", LEGAL],
            ["PP-RG-02", "Bank details.", "Bank name, account name, account number, verification flow, and privacy wording.", LEGAL],
            ["PP-RG-03", "Referral code.", "Rules for optional sponsor/referrer code during registration.", APPROVAL],
            ["PP-RG-04", "Approval state.", "Pending, approved, rejected, suspended, and reactivation wording.", APPROVAL],
            ["PP-RG-05", "Partner ID.", "Format, uniqueness rule, and when it is generated.", APPROVAL],
        ]
    )

    add_section_break(doc, "11. Contact Page Requirements")
    add_heading(doc, "Contact Page Content", 2)
    add_requirement_table(
        doc,
        [
            ["CT-01", "Contact Intro Slot.", "Approved intro copy and Contact Image Canvas asset.", REQUIRED],
            ["CT-02", "Agriculture Contact Slot.", "Phone/email/WhatsApp/routing rule for agriculture inquiries.", REQUIRED],
            ["CT-03", "Agro-Real Estate Contact Slot.", "Phone/email/WhatsApp/routing rule for Agro-Real Estate inquiries.", REQUIRED],
            ["CT-04", "Partner Program Contact Slot.", "Phone/email/WhatsApp/routing rule for partner inquiries.", REQUIRED],
            ["CT-05", "Map / Location Image Canvas.", "Exact address, map embed preference, branch display rules, and image fallback.", REQUIRED],
            ["CT-06", "Working hours.", "Opening days, hours, public holiday rules, and response-time expectations.", REQUIRED],
            ["CT-07", "Contact form privacy.", "Consent wording, data storage rules, spam protection, and response process.", LEGAL],
        ]
    )

    add_heading(doc, "Contact Form Fields", 2)
    add_table(
        doc,
        ["Field", "Current frontend", "Final requirement"],
        [
            ["Name", "Text input.", "Required field. Define validation and whether full name is required."],
            ["Email", "Email input.", "Required or optional decision. Define validation and fallback phone/WhatsApp if email is optional."],
            ["Inquiry", "Dropdown: Agriculture, Agro-Real Estate, Partner Program.", "Confirm dropdown options and internal routing for each option."],
            ["Message", "Textarea.", "Required field. Define maximum length and blocked content rules."],
            ["Submit", "Button only in current frontend.", "Define success message, error message, email notification, CRM/storage destination, and anti-spam requirement."],
        ],
        [1.2, 2.35, 4.1],
    )

    add_section_break(doc, "12. Login and Shared Identity Requirements")
    add_heading(doc, "Login Page", 2)
    add_requirement_table(
        doc,
        [
            ["LG-01", "Account Access Slot.", "Approved copy explaining portal access without exposing admin access.", REQUIRED],
            ["LG-02", "Login Image Canvas.", "Approved image or illustration asset for portal access.", MEDIA],
            ["LG-03", "Email field.", "Final validation, error, required state, and autofill behavior.", APPROVAL],
            ["LG-04", "Password field.", "Final validation, show/hide password decision, forgot-password flow, and security copy.", APPROVAL],
            ["LG-05", "Demo dashboard buttons.", "Frontend demo only: Open Partner Dashboard and Open Client Dashboard. No Admin Dashboard button.", APPROVAL],
            ["LG-06", "Shared identity model.", "Future implementation uses one user identity that can have partner and client roles.", APPROVAL],
            ["LG-07", "Role switcher.", "Future users with both roles should switch between Partner and Client dashboards after login.", APPROVAL],
        ]
    )

    add_note(
        doc,
        "Admin exclusion",
        "Admin management may exist later as a separate internal admin application or protected internal area. It must not be shown as a public navigation item or demo button on the user-facing frontend.",
        LIGHT_ORANGE,
    )

    add_section_break(doc, "13. Partner Dashboard Requirements")
    add_heading(doc, "Dashboard Purpose", 2)
    add_para(
        doc,
        "The Partner Dashboard is a user-facing portal shell for partners. It should display partner identity, tier status, referral tools, team growth, earnings, wallet, withdrawals, products, marketing resources, notifications, and settings after future backend integration.",
    )

    add_requirement_table(
        doc,
        [
            ["PD-AC-01", "Profile.", "Full Name, Passport, ID, Address, Bank Details. Define exact fields, edit permissions, KYC status, and verification state.", LEGAL],
            ["PD-AC-02", "Partner ID.", "Partner ID, Referral Code, and related canvas. Define ID format and copy/share behavior.", APPROVAL],
            ["PD-AC-03", "Partner Tier.", "Current Tier and Tier Status. Define states, upgrade path, renewal rules, and restrictions.", APPROVAL],
            ["PD-GR-01", "Referral Link.", "Referral Link and QR Code. Define generated link format, copy action, QR download rule, and tracking.", APPROVAL],
            ["PD-GR-02", "Referral Tree.", "Direct Referrals, Downline, and wide canvas. Define tree depth, display mode, privacy rules, and empty state.", LEGAL],
            ["PD-GR-03", "Team.", "Team Members and Team Activity. Define what partner can see and what is hidden for privacy.", LEGAL],
            ["PD-EA-01", "Earnings.", "Commission and Bonuses. Define pending/available/reversed states and approved labels.", LEGAL],
            ["PD-EA-02", "Wallet.", "Wallet Balance and Wallet Activity. Define source of funds, payout rules, ledger rules, and audit trail.", LEGAL],
            ["PD-EA-03", "Withdrawals.", "Withdrawal Requests and Withdrawal History. Define request form, approval flow, status labels, and bank details use.", LEGAL],
            ["PD-RS-01", "Products.", "Live Fish, Dried/Packaged Fish, Palm, Pepper. Define partner-facing catalog, sales process, and permissions.", APPROVAL],
            ["PD-RS-02", "Marketing Materials.", "Banner, Flyer, Brochure, and canvas. Define downloadable files, brand rules, and usage restrictions.", APPROVAL],
            ["PD-RS-03", "Notifications.", "Define announcement categories, unread state, detail view, and delivery channels.", APPROVAL],
            ["PD-RS-04", "Settings.", "Password and Account Preferences. Define editable preferences, security settings, and account closure request process.", APPROVAL],
        ]
    )

    add_heading(doc, "Partner Dashboard States", 2)
    add_bullets(
        doc,
        [
            "Empty state for a partner with no referrals, no sales, no earnings, and no withdrawals.",
            "Pending approval state before Partner ID becomes active.",
            "Suspended or rejected state if partner account is not eligible.",
            "Loading and error states for referral tree, wallet, marketing materials, and notifications.",
            "Mobile layout where dashboard cards remain scannable and do not overflow.",
        ],
    )

    add_section_break(doc, "14. Client Dashboard Requirements")
    add_heading(doc, "Dashboard Purpose", 2)
    add_para(
        doc,
        "The Client Dashboard is a user-facing portal shell for Agro-Real Estate clients. It should display client identity, land records, managed cultivation, investment plans, farm stage, photos, reports, yield history, notifications, and settings after future backend integration.",
    )

    add_requirement_table(
        doc,
        [
            ["CD-AC-01", "Profile.", "Full Name, Email, Phone, Address. Define edit permissions and verification requirements.", LEGAL],
            ["CD-AC-02", "Client ID.", "Client ID and Account Type. Define ID format, account types, and when ID is generated.", APPROVAL],
            ["CD-AR-01", "Agro-Real Estate.", "Buy Cultivated Farmland and Buy Land & Managed Cultivation. Define entry points and client-specific actions.", APPROVAL],
            ["CD-AR-02", "Land Records.", "Land Records and Documents. Define document types, upload/download permissions, legal verification status, and privacy.", LEGAL],
            ["CD-AR-03", "Managed Cultivation.", "Cultivation and Management. Define crop, land, service duration, management scope, and reporting cadence.", LEGAL],
            ["CD-AR-04", "Investment Plans.", "Investment Plans. Define plan enrollment, terms, payment state, renewal, and cancellation.", LEGAL],
            ["CD-FU-01", "Farm Stage.", "Farm Stage and Current Activity. Define stage names, update source, date display, and status history.", APPROVAL],
            ["CD-FU-02", "Photos.", "Farm Photos canvas. Define upload source, captions, dates, visibility, and moderation.", MEDIA],
            ["CD-FU-03", "Reports.", "Progress Reports wide card. Define report format, frequency, downloadable files, and approval process.", APPROVAL],
            ["CD-FU-04", "Yield History.", "Yield History. Define units, harvest period, sale/use of produce, and how client yield is calculated.", LEGAL],
            ["CD-ST-01", "Notifications.", "Define client notification categories and delivery channels.", APPROVAL],
            ["CD-ST-02", "Settings.", "Password and Account Preferences. Define editable preferences and support request process.", APPROVAL],
        ]
    )

    add_heading(doc, "Client Dashboard States", 2)
    add_bullets(
        doc,
        [
            "Empty state for a client with no land record yet.",
            "Pending document verification state.",
            "Managed cultivation not started state.",
            "Farm stage update available state.",
            "Report unavailable, report processing, and report ready states.",
            "Mobile layout where records, photos, reports, and yield history stay readable.",
        ],
    )

    add_section_break(doc, "15. Interactions and UX Requirements")
    add_heading(doc, "Carousels", 2)
    add_requirement_table(
        doc,
        [
            ["UX-CR-01", "Home hero carousel.", "Four slides, 7 second autoplay in current frontend, arrows, dots, and smooth horizontal movement.", APPROVAL],
            ["UX-CR-02", "Hero carousel mobile behavior.", "Slides must not clip text, CTA buttons, image canvases, or dots on small screens.", APPROVAL],
            ["UX-CR-03", "Business carousel.", "Shows Agriculture and Agro-Real Estate one at a time with arrows and dots.", APPROVAL],
            ["UX-CR-04", "Accessibility.", "Buttons require labels, keyboard focus, reduced-motion handling, and visible controls.", APPROVAL],
        ]
    )

    add_heading(doc, "Responsive Requirements", 2)
    add_bullets(
        doc,
        [
            "Header must stay usable on mobile with all public navigation items available.",
            "Hero, business carousel, product cards, dashboard cards, and forms must not overflow horizontally.",
            "Canvas placeholders and final images should be slimmer where appropriate, especially in repeated sections.",
            "CTA groups should wrap cleanly on mobile and remain visually balanced on desktop.",
            "Dashboard cards should stack logically by section and remain readable without fake numbers.",
        ],
    )

    add_heading(doc, "Form Requirements", 2)
    add_requirement_table(
        doc,
        [
            ["UX-FM-01", "Login form.", "Final validation, loading, success, error, invalid credentials, locked account, and password reset states.", APPROVAL],
            ["UX-FM-02", "Contact form.", "Required fields, validation, spam protection, success state, error state, and internal notification path.", APPROVAL],
            ["UX-FM-03", "Partner registration future form.", "KYC, ID upload, passport upload, bank details, sponsor code, terms acceptance, and approval state.", LEGAL],
            ["UX-FM-04", "Agro-Real Estate inquiry future form.", "Inquiry category, land/investment interest, contact details, preferred callback method, and consent wording.", LEGAL],
        ]
    )

    add_section_break(doc, "16. CMS and Data Requirements")
    add_heading(doc, "Future CMS Content Types", 2)
    add_table(
        doc,
        ["Content type", "Fields required", "Notes"],
        [
            ["Business", "Name, slug, status, overview, hero image, sections, gallery, SEO fields.", "Currently Agriculture and Agro-Real Estate only."],
            ["Business Section", "Business, title, slug, body, image, CTA, order, SEO fields.", "Used by Agriculture and Agro-Real Estate subpages."],
            ["Agriculture Product Category", "Name, slug, image, summary, availability, inquiry CTA, ordering rule.", "Future marketplace grows from this model."],
            ["Agro-Real Estate Investment Plan", "Name, terms, inclusions, price if applicable, risk note, documents, inquiry CTA.", "Requires legal approval."],
            ["Partner Tier", "Name, fee if any, benefits, eligibility, commission rules, status.", "Use tier naming consistently."],
            ["FAQ", "Question, answer, category, route visibility, status.", "No generic questions."],
            ["News & Update", "Title, category, image, body, date, author/source, publish status.", "Only publish real updates."],
            ["Gallery Asset", "Image, caption, alt text, owner, business, location, publish status.", "Ownership confirmation required."],
            ["Dashboard Notification", "Audience, title, message, route/action, publish date, read state.", "Future backend feature."],
        ],
        [1.8, 3.65, 2.2],
    )

    add_heading(doc, "Shared Identity Future Model", 2)
    add_bullets(
        doc,
        [
            "One user account can become a client, a partner, or both.",
            "Partner-specific data should live in partner profile records.",
            "Client-specific data should live in client/agro-real-estate records.",
            "Role-based dashboard access should show only the relevant portal to the user.",
            "Admin access should be separated from user-facing navigation and guarded by internal authorization.",
        ],
    )

    add_section_break(doc, "17. SEO, Analytics, Accessibility, and Performance")
    add_heading(doc, "SEO Requirements", 2)
    add_requirement_table(
        doc,
        [
            ["SEO-01", "Page titles.", "Unique approved title for every route.", REQUIRED],
            ["SEO-02", "Meta descriptions.", "Approved description for every public route.", REQUIRED],
            ["SEO-03", "Open Graph images.", "Approved social preview image per major page.", MEDIA],
            ["SEO-04", "Structured data.", "Organization, LocalBusiness if applicable, Product if marketplace launches, FAQ where approved.", APPROVAL],
            ["SEO-05", "Sitemap and robots.", "Include only public routes. Exclude dashboards from public indexing.", APPROVAL],
        ]
    )

    add_heading(doc, "Accessibility Requirements", 2)
    add_bullets(
        doc,
        [
            "Every image requires meaningful alt text or decorative marking.",
            "Buttons and carousel controls require accessible labels.",
            "Color contrast must be checked for orange buttons, green agriculture accents, warm backgrounds, and ink text.",
            "Forms require labels, errors, and focus states.",
            "Keyboard users must be able to navigate header, carousels, forms, and dashboards.",
        ],
    )

    add_heading(doc, "Analytics Requirements", 2)
    add_requirement_table(
        doc,
        [
            ["AN-01", "Analytics provider.", "Confirm Google Analytics, Meta Pixel, TikTok Pixel, server-side analytics, or none.", REQUIRED],
            ["AN-02", "Tracked events.", "Hero CTA clicks, business carousel clicks, contact submits, partner login/register clicks, product inquiry clicks, investment inquiry clicks.", APPROVAL],
            ["AN-03", "Consent.", "Cookie consent and privacy policy requirement if tracking tools are used.", LEGAL],
        ]
    )

    add_section_break(doc, "18. Security, Privacy, and Compliance")
    add_heading(doc, "Security and Privacy Requirements", 2)
    add_requirement_table(
        doc,
        [
            ["SP-01", "KYC data handling.", "Passport, ID, address, and bank details require secure storage, access control, retention policy, and privacy notice.", LEGAL],
            ["SP-02", "Investment documents.", "Land records and client documents require private access, audit logs, and download permissions.", LEGAL],
            ["SP-03", "Wallet and commissions.", "Future wallet ledger must be append-only, auditable, and separated from display-only dashboard totals.", LEGAL],
            ["SP-04", "Form submissions.", "Protect against spam, injection, excessive submissions, and unauthorized access.", APPROVAL],
            ["SP-05", "Admin access.", "Admin tools must require separate protected access and must not be exposed in public frontend navigation.", LEGAL],
            ["SP-06", "Terms acceptance.", "Partner, investment, managed cultivation, and marketplace terms should be accepted and logged before transactions.", LEGAL],
        ]
    )

    add_heading(doc, "Required Legal Pages or Documents", 2)
    add_bullets(
        doc,
        [
            "Privacy Policy.",
            "Website Terms of Use.",
            "Partner Program Terms.",
            "Agro-Real Estate Investment or Purchase Terms.",
            "Managed Cultivation Agreement or terms summary.",
            "Refund, cancellation, withdrawal, and dispute rules where applicable.",
            "Product ordering, delivery, pickup, and return policy when product sales go live.",
        ],
    )

    add_section_break(doc, "19. QA and Acceptance Criteria")
    add_heading(doc, "Content Acceptance", 2)
    add_bullets(
        doc,
        [
            "No lorem ipsum or generic placeholder copy remains on public pages.",
            "No fake price, product detail, business description, CEO detail, testimonial, statistic, location, ROI, yield, commission, or partner benefit is published.",
            "Every image canvas has been replaced with an approved image or intentionally left as a labelled structural placeholder for internal review only.",
            "Every CTA target is confirmed and works.",
            "Real Estate, Construction, African Tales, and Admin Dashboard remain excluded from user-facing navigation.",
        ],
    )

    add_heading(doc, "Technical Acceptance", 2)
    add_bullets(
        doc,
        [
            "All public routes load successfully.",
            "Login demo buttons open Partner Dashboard and Client Dashboard only.",
            "Admin dashboard route is not generated as a public frontend route.",
            "Desktop, tablet, and mobile layouts show no clipping, overlap, or horizontal overflow.",
            "Carousels behave correctly on desktop and mobile.",
            "Forms have clear fields, validation states, and final submit behavior after backend integration.",
            "Dashboards remain visually consistent and use real data only after integration.",
        ],
    )

    add_heading(doc, "Pre-Launch Route Checklist", 2)
    add_table(
        doc,
        ["Route", "Desktop", "Mobile", "Content approved", "Images approved", "CTA tested"],
        [
            ["/", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/about", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/businesses", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/businesses/agriculture", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/businesses/agriculture/crop-farming", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/businesses/agriculture/livestock", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/businesses/agriculture/livestock/poultry", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/businesses/agriculture/livestock/fishery", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/businesses/agriculture/livestock/bsf", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/businesses/agriculture/products", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/businesses/agriculture/gallery", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/businesses/agro-real-estate", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/businesses/agro-real-estate/buy-cultivated-farmland", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/businesses/agro-real-estate/managed-cultivation", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/businesses/agro-real-estate/investment-plans", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/businesses/agro-real-estate/faqs", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/businesses/agro-real-estate/gallery", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/partner-program", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/contact", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/login", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/dashboard/partner", "Pending", "Pending", "Pending", "Pending", "Pending"],
            ["/dashboard/client", "Pending", "Pending", "Pending", "Pending", "Pending"],
        ],
        [2.35, 0.95, 0.95, 1.4, 1.4, 1.25],
    )

    add_section_break(doc, "20. Content Handover Checklist")
    add_heading(doc, "Client Content Package", 2)
    add_requirement_table(
        doc,
        [
            ["CH-01", "Company profile pack.", "Official name, about, vision, mission, why choose us, CEO profile, logo, brand colors.", REQUIRED],
            ["CH-02", "Agriculture pack.", "Overview, crop farming, livestock, product categories, gallery images, contact process.", REQUIRED],
            ["CH-03", "Agro-Real Estate pack.", "Overview, land offers, managed cultivation, investment plans, FAQs, legal terms, gallery images.", LEGAL],
            ["CH-04", "Partner Program pack.", "Overview, benefits, tiers, compensation, registration rules, KYC rules, terms, FAQs.", LEGAL],
            ["CH-05", "Contact pack.", "Phone, email, WhatsApp, address, map, branches, working hours, inquiry routing.", REQUIRED],
            ["CH-06", "Image pack.", "Named folders matching each page and canvas, with captions and alt text.", MEDIA],
            ["CH-07", "Legal pack.", "Privacy, terms, partner terms, investment terms, refund/cancellation, disclaimers.", LEGAL],
            ["CH-08", "Approval pack.", "Final sign-off from De Jadon Group for all content before public launch.", APPROVAL],
        ]
    )

    add_heading(doc, "Recommended File Naming", 2)
    add_table(
        doc,
        ["Asset group", "Suggested folder", "Example file name"],
        [
            ["Logo", "brand", "de-jadon-group-logo-primary.svg"],
            ["Home hero", "home/hero", "home-hero-agriculture-01.jpg"],
            ["About", "about", "about-de-jadon-group-01.jpg"],
            ["CEO", "about/ceo", "ceo-portrait-01.jpg"],
            ["Agriculture", "businesses/agriculture", "agriculture-overview-01.jpg"],
            ["Products", "businesses/agriculture/products", "product-live-fish-01.jpg"],
            ["Agro-Real Estate", "businesses/agro-real-estate", "agro-real-estate-overview-01.jpg"],
            ["Partner Program", "partner-program", "partner-program-hero-01.jpg"],
            ["Contact", "contact", "contact-location-map-01.jpg"],
        ],
        [1.7, 2.6, 3.35],
    )

    add_section_break(doc, "21. Final Sign-Off")
    add_heading(doc, "Approval Summary", 2)
    add_table(
        doc,
        ["Approval area", "Approver name", "Signature", "Date", "Status"],
        [
            ["Company and brand content", "", "", "", "Pending"],
            ["Agriculture content", "", "", "", "Pending"],
            ["Agro-Real Estate content and terms", "", "", "", "Pending"],
            ["Partner Program content and terms", "", "", "", "Pending"],
            ["Contact and operations", "", "", "", "Pending"],
            ["Legal and privacy", "", "", "", "Pending"],
            ["Final website launch", "", "", "", "Pending"],
        ],
        [2.2, 1.8, 1.4, 1.1, 1.15],
    )

    add_note(
        doc,
        "Final rule",
        "The website should launch only after all required information has been supplied, reviewed, and approved. The frontend skeleton is intentionally incomplete until that content is provided.",
        LIGHT_GREEN,
    )

    return doc


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
